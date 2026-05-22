from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any


@dataclass(frozen=True)
class ImportedReference:
    patches: list[dict[str, object]]
    source: dict[str, object]

    @property
    def color_map(self) -> dict[str, dict[str, list[float]]]:
        return {
            str(patch["label"]): {
                "sRGB": list(patch["sRGB"]),
                "Lab_d50": list(patch["Lab_d50"]),
                "Lab_d65": list(patch["Lab_d65"]),
            }
            for patch in self.patches
        }


def _header_text(value: object) -> str:
    text = str(value or "").strip().lower()
    for old, new in {
        "（": "(",
        "）": ")",
        "＊": "*",
        "　": " ",
        "，": ",",
    }.items():
        text = text.replace(old, new)
    return text


def _header_tokens(value: object) -> tuple[str, ...]:
    return tuple(
        token
        for token in re.split(r"[^0-9a-zA-Z#\u4e00-\u9fff]+", _header_text(value))
        if token
    )


def _normalize_header(value: object) -> str:
    return "".join(_header_tokens(value))


def _find_metadata_column(headers: list[tuple[str, tuple[str, ...]]], candidates: set[str]) -> int | None:
    for idx, (compact, tokens) in enumerate(headers):
        if compact in candidates or any(token in candidates for token in tokens):
            return idx
    return None


def _leaf_token(tokens: tuple[str, ...]) -> str:
    return tokens[-1] if tokens else ""


def _matches_channel_leaf(token: str, channel: str, aliases: dict[str, set[str]]) -> bool:
    return token in aliases[channel]


def _find_srgb_column(headers: list[tuple[str, tuple[str, ...]]], channel: str) -> int | None:
    aliases = {
        "r": {"r", "red", "红"},
        "g": {"g", "green", "绿"},
        "b": {"b", "blue", "蓝"},
    }
    for idx, header in enumerate(headers):
        compact, tokens = header
        leaf = _leaf_token(tokens)
        if _matches_channel_leaf(leaf, channel, aliases) and (
            "srgb" in tokens
            or "rgb" in tokens
            or compact in {f"srgb{channel}", f"rgb{channel}"}
        ):
            return idx
        if compact in aliases[channel]:
            return idx
    return None


def _find_lab_column(headers: list[tuple[str, tuple[str, ...]]], white_point: str, channel: str) -> int | None:
    aliases = {
        "l": {"l", "l*", "明度"},
        "a": {"a", "a*"},
        "b": {"b", "b*"},
    }
    white = white_point.lower()
    for idx, (compact, tokens) in enumerate(headers):
        leaf = _leaf_token(tokens)
        if not _matches_channel_leaf(leaf, channel, aliases):
            continue
        if white not in tokens and white not in compact:
            continue
        if "lab" in tokens or "cie" in tokens or "lab" in compact or white in tokens:
            return idx
    return None


def _header_mapping(row: tuple[object, ...]) -> dict[str, int] | None:
    headers = [(_normalize_header(cell), _header_tokens(cell)) for cell in row]
    mapping: dict[str, int | None] = {
        "index": _find_metadata_column(headers, {"编号", "序号", "patchno", "patchnumber", "no", "#"}),
        "label": _find_metadata_column(headers, {"名称", "name", "label", "色块", "number"}),
        "srgb_r": _find_srgb_column(headers, "r"),
        "srgb_g": _find_srgb_column(headers, "g"),
        "srgb_b": _find_srgb_column(headers, "b"),
        "lab_d50_l": _find_lab_column(headers, "D50", "l"),
        "lab_d50_a": _find_lab_column(headers, "D50", "a"),
        "lab_d50_b": _find_lab_column(headers, "D50", "b"),
        "lab_d65_l": _find_lab_column(headers, "D65", "l"),
        "lab_d65_a": _find_lab_column(headers, "D65", "a"),
        "lab_d65_b": _find_lab_column(headers, "D65", "b"),
    }
    required = [
        "index",
        "srgb_r",
        "srgb_g",
        "srgb_b",
        "lab_d50_l",
        "lab_d50_a",
        "lab_d50_b",
        "lab_d65_l",
        "lab_d65_a",
        "lab_d65_b",
    ]
    if any(mapping[key] is None for key in required):
        return None
    required_indices = [mapping[key] for key in required]
    if len(set(required_indices)) != len(required_indices):
        return None
    return {key: int(value) for key, value in mapping.items() if value is not None}


def _combined_header_row(parent_row: tuple[object, ...], child_row: tuple[object, ...]) -> tuple[str, ...]:
    width = max(len(parent_row), len(child_row))
    combined: list[str] = []
    current_parent = ""
    for idx in range(width):
        parent = parent_row[idx] if idx < len(parent_row) else None
        child = child_row[idx] if idx < len(child_row) else None
        if parent not in {None, ""}:
            current_parent = str(parent)
        if child not in {None, ""}:
            combined.append(f"{current_parent} {child}".strip())
        else:
            combined.append(current_parent)
    return tuple(combined)


def _as_number(value: object, field: str) -> float:
    if value is None or value == "":
        raise ValueError(f"参考色值缺少字段：{field}")
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"参考色值字段 {field} 不是数字：{value!r}") from exc


def _validate_srgb(values: list[float], index: int) -> list[int]:
    out: list[int] = []
    for value in values:
        if not 0.0 <= value <= 255.0:
            raise ValueError(f"第 {index} 行 sRGB 超出 0~255：{values}")
        out.append(int(round(value)))
    return out


def _validate_lab(values: list[float], index: int, white: str) -> list[float]:
    L, a, b = values
    if not 0.0 <= L <= 100.0:
        raise ValueError(f"第 {index} 行 Lab {white} 的 L* 超出 0~100：{values}")
    if not -150.0 <= a <= 150.0 or not -150.0 <= b <= 150.0:
        raise ValueError(f"第 {index} 行 Lab {white} 的 a*/b* 超出合理范围：{values}")
    return [float(L), float(a), float(b)]


def _serial_from_name(file_name: str) -> str | None:
    match = re.search(r"(\d{5,})", Path(file_name or "").stem)
    return match.group(1) if match else None


def _find_reference_table(workbook: object) -> tuple[list[tuple[object, ...]], dict[str, int], int]:
    sheetnames = list(getattr(workbook, "sheetnames", []) or [])
    if "Sheet1" in sheetnames:
        sheetnames = ["Sheet1", *[name for name in sheetnames if name != "Sheet1"]]
    for sheet_name in sheetnames:
        sheet = workbook[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        for row_index, row in enumerate(rows[:20]):
            mapping = _header_mapping(row)
            if mapping is not None:
                return rows, mapping, row_index + 1
            if row_index + 1 < len(rows):
                mapping = _header_mapping(_combined_header_row(row, rows[row_index + 1]))
                if mapping is not None:
                    return rows, mapping, row_index + 2
    raise ValueError("xlsx 参考文件缺少编号、sRGB R/G/B、Lab D50、Lab D65 必需列。")


def parse_reference_xlsx(raw: bytes | str | Path, *, file_name: str = "") -> ImportedReference:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ValueError("解析 xlsx 参考文件需要 openpyxl。") from exc

    if isinstance(raw, (str, Path)):
        path = Path(raw)
        data: bytes | str | Path = path
        source_name = file_name or path.name
    else:
        data = io.BytesIO(raw)
        source_name = file_name or "uploaded_reference.xlsx"

    try:
        workbook = load_workbook(data, data_only=True, read_only=True)
    except (OSError, ValueError, TypeError, zipfile.BadZipFile, InvalidFileException) as exc:
        raise ValueError(f"xlsx 参考文件无法解析：{exc}") from exc
    rows, mapping, data_start_index = _find_reference_table(workbook)

    patches: list[dict[str, object]] = []
    for row in rows[data_start_index:]:
        if row is None or all(cell in {None, ""} for cell in row):
            continue
        if row[mapping["index"]] in {None, ""}:
            continue
        idx = int(round(_as_number(row[mapping["index"]], "编号")))
        label = str(row[mapping["label"]]).strip() if "label" in mapping and row[mapping["label"]] not in {None, ""} else f"Patch {idx}"
        srgb = _validate_srgb(
            [
                _as_number(row[mapping["srgb_r"]], "sRGB R"),
                _as_number(row[mapping["srgb_g"]], "sRGB G"),
                _as_number(row[mapping["srgb_b"]], "sRGB B"),
            ],
            idx,
        )
        lab_d50 = _validate_lab(
            [
                _as_number(row[mapping["lab_d50_l"]], "Lab D50 L*"),
                _as_number(row[mapping["lab_d50_a"]], "Lab D50 a*"),
                _as_number(row[mapping["lab_d50_b"]], "Lab D50 b*"),
            ],
            idx,
            "D50",
        )
        lab_d65 = _validate_lab(
            [
                _as_number(row[mapping["lab_d65_l"]], "Lab D65 L*"),
                _as_number(row[mapping["lab_d65_a"]], "Lab D65 a*"),
                _as_number(row[mapping["lab_d65_b"]], "Lab D65 b*"),
            ],
            idx,
            "D65",
        )
        patches.append(
            {
                "index": idx,
                "label": label,
                "sRGB": srgb,
                "Lab_d50": lab_d50,
                "Lab_d65": lab_d65,
            }
        )

    if len(patches) != 24:
        raise ValueError(f"xlsx 参考文件必须包含 24 个色块，当前为 {len(patches)} 个。")
    patches.sort(key=lambda item: int(item["index"]))
    indices = [int(patch["index"]) for patch in patches]
    if indices != list(range(1, 25)):
        raise ValueError(f"xlsx 参考文件色块编号必须连续为 1~24，当前为 {indices}。")
    labels = [str(patch["label"]) for patch in patches]
    duplicate_labels = sorted({label for label in labels if labels.count(label) > 1})
    if duplicate_labels:
        raise ValueError(f"xlsx 参考文件色块名称必须唯一，重复名称：{duplicate_labels}。")
    source = {
        "source_type": "imported_xlsx",
        "source_name": source_name,
        "serial_number": _serial_from_name(source_name),
        "patch_count": 24,
        "lab_white_point": "D50",
    }
    return ImportedReference(patches=patches, source=source)


def reference_summary(reference: ImportedReference, token: str | None = None) -> dict[str, Any]:
    return {
        "referenceToken": token,
        "source": reference.source,
        "patches": reference.patches,
        "patchCount": len(reference.patches),
    }
